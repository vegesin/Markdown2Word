"""Regenerate the small Word requirements fixture from its Markdown source."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


def main():
    root=Path(__file__).parent/'fixtures'
    doc=Document()
    section=doc.sections[0]
    section.page_width,section.page_height=Mm(210),Mm(297)
    for side in ('top','bottom','left','right'):
        setattr(section,side+'_margin',Mm(25))
    for name,size in [('Normal',11),('Title',20),('Heading 1',14)]:
        style=doc.styles[name]
        style.font.name='Times New Roman'
        style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.add_paragraph('技术报告格式要求',style='Title')
    doc.add_paragraph('本文件用于测试 Markdown2Word 的格式要求提取与模板生成。以下段落和表格规定正文、标题及图表的排版要求。')
    lines=(root/'requirements.md').read_text(encoding='utf-8').splitlines()
    for line in lines[:6]:
        doc.add_paragraph(line)
    doc.add_paragraph('图表与代码格式',style='Heading 1')
    table=doc.add_table(rows=0,cols=2)
    table.style='Table Grid'
    for line in lines[6:]:
        label,body=line.split('：',1)
        cells=table.add_row().cells
        cells[0].text=label
        cells[1].text=body
    doc.save(root/'requirements.docx')


if __name__=='__main__':
    main()
