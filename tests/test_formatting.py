from __future__ import annotations

import copy
import base64
import io
import json
import sys
import threading
import tempfile
import unittest
import urllib.error
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from docx.oxml.ns import qn

from core.pipeline import ConversionConfig, PipelineRunner, find_pandoc, validate
from formatting.spec import resolve_spec, load_spec
from formatting.word import build_project, audit_package, configure_document, apply_content
from requirements.extract import extract
from requirements.parser import local_parse, validate_draft, decode_draft
from requirements.ai import chat, NoRedirect
from gui.bridge import DesktopApi

FIXTURES=Path(__file__).parent/'fixtures'


class FormattingTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.basic={'schema_version':'1.0','name':'测试规范'}

    def test_schema_rejects_unknown_unsupported_and_invalid_units(self):
        for update in [{'schema_version':'2.0'}, {'unknown':1}, {'styles':{'body':{'font':{'size_pt':True}}}},
                       {'styles':{'body':{'paragraph':{'line_spacing':{'mode':'exact','value':0}}}}},
                       {'captions':{'figure':{'scope':'chapter'}}}, {'page':{'margins_mm':{'left':float('nan')}}}]:
            with self.subTest(update=update), self.assertRaises(ValueError):
                resolve_spec({**self.basic,**update})

    def test_inheritance_cycles_spacing_and_page_geometry(self):
        with self.assertRaisesRegex(ValueError,'循环'):
            resolve_spec({**self.basic,'styles':{'body':{'based_on':'heading_1'}}})
        with self.assertRaisesRegex(ValueError,'正文区域'):
            resolve_spec({**self.basic,'page':{'margins_mm':{'left':100,'right':100}}})
        spec=resolve_spec({**self.basic,'styles':{'body':{'font':{'size_pt':11},'paragraph':{'line_spacing':{'mode':'single'}}}}})
        self.assertEqual(spec['styles']['first_paragraph']['font']['size_pt'],11)
        self.assertEqual(spec['styles']['body']['paragraph']['line_spacing'],{'mode':'single'})
        self.assertEqual(resolve_spec(spec),spec)

    def test_local_parse_fixture_and_conflicts(self):
        blocks=extract(FIXTURES/'requirements.md')
        draft=local_parse(blocks)
        self.assertEqual(draft['issues'],[])
        self.assertEqual(draft['spec']['page']['margins_mm']['left'],30)
        self.assertEqual(draft['spec']['styles']['body']['font']['size_pt'],12)
        validate_draft(draft,blocks)
        conflict=local_parse(extract(FIXTURES/'conflicting-requirements.md'))
        self.assertTrue(any(i['kind']=='conflict' for i in conflict['issues']))
        self.assertTrue(any(i['kind']=='unresolved' for i in conflict['issues']))

    def test_word_extraction_preserves_table_order_without_word(self):
        doc=Document()
        doc.add_paragraph('正文：宋体，小四')
        table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='表头'
        table.cell(0,1).text='黑体，五号'
        doc.add_paragraph('一级标题：黑体')
        path=self.root/'requirements.docx'
        doc.save(path)
        blocks=extract(path)
        self.assertEqual([b['text'] for b in blocks],['正文：宋体，小四','表头；黑体，五号','一级标题：黑体'])
        self.assertIn('表格行',blocks[1]['source'])

    def test_bundled_word_fixture_contains_requirements(self):
        blocks=extract(FIXTURES/'requirements.docx')
        self.assertTrue(any('正文：宋体' in b['text'] for b in blocks))
        self.assertTrue(any('表格行' in b['source'] for b in blocks))
        result=local_parse(blocks)
        self.assertEqual(result['spec']['styles']['body']['font']['size_pt'],12)

    def test_project_generation_is_clean_and_no_clobber(self):
        output=build_project(self.basic,self.root/'project')
        expected={'spec.json','resolved-spec.json','reference.docx','preview.docx','report.json','manifest.json'}
        self.assertEqual({p.name for p in output.iterdir()},expected)
        audit_package(output/'reference.docx',clean_template=True)
        doc=Document(output/'reference.docx')
        self.assertEqual(doc.styles['Normal'].font.size.pt,12)
        self.assertEqual(doc.styles['Normal'].element.rPr.rFonts.get(qn('w:eastAsia')),'宋体')
        with self.assertRaises(ValueError): build_project(self.basic,output)
        with self.assertRaises(ValueError): build_project(self.basic,self.root/'cancelled',cancelled=lambda:True)
        self.assertFalse((self.root/'cancelled').exists())

    def test_template_package_rejects_active_parts(self):
        path=self.root/'bad.docx'
        with zipfile.ZipFile(path,'w') as z: z.writestr('word/vbaProject.bin',b'not executed')
        with self.assertRaises(ValueError): audit_package(path,clean_template=True)

    def test_figure_and_table_captions_have_independent_protocol_styles(self):
        from backends.stage1.build_stage1_docx import apply_captions
        spec=resolve_spec({**self.basic,'styles':{'figure_caption':{'font':{'size_pt':14},'paragraph':{'alignment':'right'}},'table_caption':{'font':{'size_pt':9}}},'captions':{'figure':{'label':'Figure'}}})
        doc=Document()
        configure_document(doc,spec)
        png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jM1sAAAAASUVORK5CYII=')
        doc.add_picture(io.BytesIO(png))
        doc.add_table(rows=1,cols=1).cell(0,0).text='header'
        apply_captions(doc,spec)
        apply_content(doc,spec)
        path=self.root/'captions.docx'
        doc.save(path)
        reopened=Document(path)
        figure=next(p for p in reopened.paragraphs if p.style.name=='Figure Caption')
        table=next(p for p in reopened.paragraphs if p.style.name=='Table Caption')
        self.assertTrue(figure.text.startswith('Figure'))
        self.assertEqual(figure.style.font.size.pt,14)
        self.assertEqual(figure.alignment,2)
        self.assertEqual(table.style.font.size.pt,9)
        self.assertTrue(figure._p.xpath('./w:fldSimple[contains(@w:instr,"SEQ")]'))
        image=next(p for p in reopened.paragraphs if p._p.xpath('.//w:drawing'))
        self.assertTrue(image.paragraph_format.keep_with_next)

    def test_malformed_word_is_reported_without_raw_traceback(self):
        path=self.root/'broken.docx'
        path.write_bytes(b'not a zip')
        with self.assertRaisesRegex(ValueError,'DOCX'): extract(path)

    def test_ai_payload_history_and_source_validation(self):
        blocks=[{'source':'行 1','text':'正文：宋体，小四'}]
        draft=local_parse(blocks)
        response=io.BytesIO(json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(draft,ensure_ascii=False)}}]}).encode())
        opener=Mock()
        opener.open.return_value=response
        result,history=chat({'base_url':'https://service.example/v1','model':'example'},blocks,'请生成',api_key='secret-test',opener=opener)
        request=opener.open.call_args.args[0]
        payload=json.loads(request.data)
        self.assertEqual(request.full_url,'https://service.example/v1/chat/completions')
        self.assertIn('JSON Schema',payload['messages'][0]['content'])
        self.assertNotIn('secret-test',json.dumps(payload))
        self.assertEqual(result,draft)
        self.assertEqual(len(history),2)
        draft['evidence'][0]['quote']='伪造来源'
        with self.assertRaisesRegex(ValueError,'来源'): validate_draft(draft,blocks)

    def test_ai_network_errors_and_truncation_do_not_expose_secret(self):
        config={'base_url':'https://service.example/v1','model':'test'}
        opener=Mock()
        opener.open.side_effect=urllib.error.HTTPError('https://x',401,'secret-test',{},None)
        with self.assertRaises(ValueError) as error: chat(config,[],'hello',api_key='secret-test',opener=opener)
        self.assertNotIn('secret-test',str(error.exception))
        self.assertIn('401',str(error.exception))
        for url in ('http://remote.example/v1','https://key:secret@remote.example/v1','file:///tmp/test'):
            with self.assertRaises(ValueError): chat({**config,'base_url':url},[],'hello',opener=opener)
        opener.open.side_effect=None
        opener.open.return_value=io.BytesIO(b'{"choices":[{"finish_reason":"length"}]}')
        with self.assertRaisesRegex(ValueError,'正常结束'): chat(config,[],'hello',opener=opener)

    def test_ai_invalid_json_and_redirect(self):
        with self.assertRaises(ValueError): decode_draft('not json')
        with self.assertRaises(ValueError): NoRedirect().redirect_request(None,None,302,'',{},'https://other.example')

    def test_gui_generation_and_api_key_not_persisted(self):
        api=DesktopApi(settings_path=self.root/'settings.json')
        api.save_settings({'ai_model':'test','api_key':'secret','ai-key':'secret'})
        self.assertNotIn('secret',(self.root/'settings.json').read_text())
        self.assertTrue(api.import_requirements(str(FIXTURES/'requirements.md'))['ok'])
        api._worker.join(10)
        draft=api.template_snapshot()['draft']
        self.assertFalse(api.generate_template(draft,str(self.root/'project'),False)['ok'])
        self.assertTrue(api.generate_template(draft,str(self.root/'project'),True)['ok'])
        api._worker.join(10)
        self.assertEqual(api.template_snapshot()['error'],'')
        self.assertTrue((self.root/'project'/'reference.docx').is_file())
        self.assertTrue(api.save_template_draft(draft,str(self.root/'draft.json'))['ok'])
        self.assertFalse(api.save_template_draft(draft,str(self.root/'draft.json'))['ok'])

    def test_gui_ai_cancellation_discards_result_and_prevents_parallel_job(self):
        api=DesktopApi(settings_path=self.root/'settings.json')
        ready,release=threading.Event(),threading.Event()
        draft={'spec':self.basic,'issues':[],'evidence':[]}
        def fake_chat(*args):
            ready.set();release.wait(3)
            return draft,[{'role':'assistant','content':'new'}]
        with patch('requirements.ai.chat',side_effect=fake_chat):
            api.template_chat({'draft':json.dumps(draft),'message':'hello','config':{}})
            self.assertTrue(ready.wait(2))
            self.assertFalse(api.import_requirements(str(FIXTURES/'requirements.md'))['ok'])
            self.assertTrue(api.cancel_template()['ok'])
            release.set();api._worker.join(5)
        self.assertEqual(api.template_snapshot()['revision'],0)
        self.assertEqual(api._chat_history,[])

    def test_protocol_and_external_template_are_exclusive(self):
        with self.assertRaisesRegex(ValueError,'不能同时'):
            validate(ConversionConfig(markdown=FIXTURES/'report.md',output=self.root/'out.docx',template=FIXTURES/'requirements.docx',format_spec=FIXTURES/'report-spec.json'))

    def test_numbering_requires_compatible_backend(self):
        path=self.root/'spec.json'
        path.write_text(json.dumps({**self.basic,'equations':{'numbering':'Simple'}}),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'MathType'):
            validate(ConversionConfig(markdown=FIXTURES/'report.md',output=self.root/'out.docx',format_spec=path,convert_formulas=False))

    def test_actual_pandoc_two_specs_control_final_document(self):
        pandoc=find_pandoc()
        if not pandoc.is_file(): self.skipTest('Pandoc is not installed')
        for i,(size,borders) in enumerate([(10.5,'grid'),(14,'none')]):
            spec={**self.basic,'styles':{'body':{'font':{'size_pt':size}},'table_body':{'font':{'size_pt':size}},'table_header':{'font':{'size_pt':size}}},'table':{'borders':borders},'toc':{'enabled':True},'headings':{'numbering':'decimal'},'equations':{'inline_min_pt':28,'display_min_pt':35}}
            path=self.root/f'spec-{i}.json'
            path.write_text(json.dumps(spec),encoding='utf-8')
            output=self.root/f'output-{i}.docx'
            runner=PipelineRunner(log_directory=self.root/'logs')
            config=ConversionConfig(markdown=FIXTURES/'report.md',output=output,format_spec=path,python=Path(sys.executable),pandoc=pandoc,convert_formulas=False,add_numbers=False)
            self.assertEqual(runner.run(config,lambda _:None),0,runner.last_error)
            doc=Document(output)
            self.assertEqual(doc.styles['Normal'].font.size.pt,size)
            self.assertEqual(doc.tables[0].cell(1,0).paragraphs[0].runs[0].font.size.pt,size)
            self.assertTrue(doc.element.xpath('.//w:fldSimple[contains(@w:instr,"TOC")]'))
            self.assertEqual(len(doc.element.xpath('.//m:oMath')),2)
            self.assertTrue(any(p._p.xpath('./w:pPr/w:numPr') for p in doc.paragraphs if p.style.name=='Heading 1'))
            for p in doc.paragraphs:
                if p._p.xpath('.//m:oMathPara'):
                    self.assertEqual(p._p.pPr.find(qn('w:spacing')).get(qn('w:line')),'700')


if __name__=='__main__':
    unittest.main()
