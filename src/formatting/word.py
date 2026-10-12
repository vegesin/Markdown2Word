from __future__ import annotations

import hashlib
import json
import os
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

from .spec import STYLE_NAMES, dump_json, resolve_spec

ALIGN = dict(left=WD_ALIGN_PARAGRAPH.LEFT, center=WD_ALIGN_PARAGRAPH.CENTER,
             right=WD_ALIGN_PARAGRAPH.RIGHT, justify=WD_ALIGN_PARAGRAPH.JUSTIFY)


def element(tag, **attributes):
    node = OxmlElement('w:' + tag)
    for key, value in attributes.items():
        node.set(qn('w:' + key), str(value))
    return node


def font_format(target, data):
    target.font.name = data['latin']
    target.font.size = Pt(data['size_pt'])
    target.font.bold = data['bold']
    target.font.italic = data['italic']
    target.font.color.rgb = RGBColor.from_string(data['color'])
    owner = target.element if hasattr(target, 'element') else target._r
    rpr = owner.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = element('rFonts')
        rpr.insert(0, fonts)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)
    for key, value in [('ascii', data['latin']), ('hAnsi', data['latin']), ('eastAsia',data['east_asia'])]:
        fonts.set(qn('w:' + key), value)
    color = rpr.find(qn('w:color'))
    if color is not None:
        for key in ('themeColor', 'themeTint', 'themeShade'):
            color.attrib.pop(qn('w:' + key), None)


def paragraph_format(target, data):
    fmt = target.paragraph_format
    fmt.alignment = ALIGN[data['alignment']]
    for key in ('keep_with_next','keep_together','page_break_before'):
        setattr(fmt, key, data[key])
    fmt.space_before = Pt(data['space_before_pt'])
    fmt.space_after = Pt(data['space_after_pt'])
    indent = data['first_line_indent']
    fmt.first_line_indent = Pt(indent['value'] if indent['unit'] == 'pt' else 0)
    owner = target.element if hasattr(target, 'element') else target._p
    ind = owner.get_or_add_pPr().find(qn('w:ind'))
    for key in ('firstLineChars','hangingChars'):
        ind.attrib.pop(qn('w:' + key), None)
    if indent['unit'] == 'char':
        ind.set(qn('w:firstLineChars' if indent['value'] >= 0 else 'w:hangingChars'), str(round(abs(indent['value']) * 100)))
    spacing = data['line_spacing']
    if spacing['mode'] in ('single','multiple'):
        fmt.line_spacing = 1.0 if spacing['mode'] == 'single' else float(spacing['value'])
    else:
        fmt.line_spacing = Pt(spacing['value'])
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY if spacing['mode'] == 'exact' else WD_LINE_SPACING.AT_LEAST


def configure_document(doc, spec):
    for role, definition in spec['styles'].items():
        name = STYLE_NAMES[role]
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        font_format(style, definition['font'])
        paragraph_format(style, definition['paragraph'])
        if role.startswith('heading_'):
            ppr = style.element.get_or_add_pPr()
            for node in list(ppr.findall(qn('w:outlineLvl'))):
                ppr.remove(node)
            ppr.append(element('outlineLvl', val=int(role.split('_')[1])-1))
    for name in ('Body Text', 'Compact'):
        if name in doc.styles:
            font_format(doc.styles[name], spec['styles']['body']['font'])
            paragraph_format(doc.styles[name], spec['styles']['body']['paragraph'])
    page = spec['page']
    width, height = (210,297) if page['size'] == 'A4' else (215.9,279.4)
    if page['orientation'] == 'landscape':
        width, height = height, width
    for section in doc.sections:
        section.page_width, section.page_height = Mm(width), Mm(height)
        from docx.enum.section import WD_ORIENT
        section.orientation = WD_ORIENT.LANDSCAPE if page['orientation'] == 'landscape' else WD_ORIENT.PORTRAIT
        for side, value in page['margins_mm'].items():
            setattr(section, side + '_margin', Mm(value))
        section.header.paragraphs[0].text = spec['header_footer']['header_text']
        p = section.footer.paragraphs[0]
        p.clear()
        if spec['header_footer']['page_number']:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(p, 'PAGE', '1')
    apply_numbering(doc, spec)


def add_field(paragraph, instruction, cached=''):
    field = element('fldSimple', instr=instruction)
    run = element('r')
    text = element('t')
    text.text = cached
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def apply_numbering(doc, spec):
    if spec['headings']['numbering'] == 'none':
        return
    root = doc.part.numbering_part.element
    existing = [node for node in root.findall(qn('w:abstractNum'))
                if node.find(qn('w:name')) is not None and node.find(qn('w:name')).get(qn('w:val')) == 'M2WHeadings']
    if existing:
        abstract_id = existing[0].get(qn('w:abstractNumId'))
        numbers = [node for node in root.findall(qn('w:num')) if node.find(qn('w:abstractNumId')).get(qn('w:val')) == abstract_id]
        if numbers:
            _bind_heading_numbering(doc, spec, int(numbers[0].get(qn('w:numId'))))
            return
    abstract_id = max([int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    num_id = max([int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))] + [0]) + 1
    abstract = element('abstractNum', abstractNumId=abstract_id)
    abstract.append(element('multiLevelType', val='multilevel'))
    abstract.append(element('name',val='M2WHeadings'))
    for i in range(spec['headings']['depth']):
        level = element('lvl', ilvl=i)
        for child in (element('start',val=1), element('numFmt',val='decimal'),
                      element('pStyle',val=doc.styles[f'Heading {i+1}'].style_id),
                      element('lvlText',val='.'.join('%'+str(j+1) for j in range(i+1))),
                      element('suff',val='space')):
            level.append(child)
        abstract.append(level)
    # abstractNum precedes num in the OOXML sequence.
    first_num = root.find(qn('w:num'))
    root.insert(list(root).index(first_num) if first_num is not None else len(root), abstract)
    num = element('num', numId=num_id)
    num.append(element('abstractNumId',val=abstract_id))
    root.append(num)
    _bind_heading_numbering(doc, spec, num_id)


def _bind_heading_numbering(doc, spec, num_id):
    for i in range(spec['headings']['depth']):
        numpr = doc.styles[f'Heading {i+1}'].element.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = i
        numpr.get_or_add_numId().val = num_id
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading '):
            level = int(p.style.name.split()[-1]) - 1
            if level < spec['headings']['depth']:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = level
                numpr.get_or_add_numId().val = num_id


def apply_content(doc, spec):
    """Apply semantic formatting after Pandoc and before Office backends."""
    from backends.stage1.build_stage1_docx import set_table_borders, set_cell_margins, set_line_spacing_at_least
    from backends.stage1.markdown_ast import content_paragraphs, formula_literals
    configure_document(doc, spec)
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        for node in list(tblpr.findall(qn('w:tblBorders'))):
            tblpr.remove(node)
        if spec['table']['borders'] == 'three_line':
            set_table_borders(table)
        else:
            borders = element('tblBorders')
            for edge in ('top','bottom','left','right','insideH','insideV'):
                borders.append(element(edge,val='single' if spec['table']['borders']=='grid' else 'nil',sz=6,color='000000'))
            tblpr.append(borders)
        for i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            for node in list(trpr.findall(qn('w:tblHeader'))):
                trpr.remove(node)
            if i == 0 and spec['table']['repeat_header']:
                trpr.append(element('tblHeader',val=1))
            for cell in row.cells:
                margin = round(spec['table']['cell_margin_mm'] * 1440 / 25.4)
                set_cell_margins(cell,margin,margin,margin,margin)
                role = 'table_header' if i == 0 else 'table_body'
                for p in cell.paragraphs:
                    p.style = doc.styles[STYLE_NAMES[role]]
                    paragraph_format(p,spec['styles'][role]['paragraph'])
                    for run in p.runs:
                        font_format(run,spec['styles'][role]['font'])
    role_by_name = {v:k for k,v in STYLE_NAMES.items()}
    for p in content_paragraphs(doc):
        role = role_by_name.get(p.style.name)
        if role:
            paragraph_format(p,spec['styles'][role]['paragraph'])
            if role == 'code':
                for run in p.runs:
                    font_format(run,spec['styles'][role]['font'])
        if p._p.xpath('.//w:drawing') and spec['captions']['figure']['enabled']:
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        literals = list(formula_literals(p))
        display = any(t.startswith('\\[') for t in literals) or bool(p._p.xpath('.//m:oMathPara'))
        inline = bool(literals) or bool(p._p.xpath('.//m:oMath'))
        if display:
            p.style = doc.styles[STYLE_NAMES['equation']]
            paragraph_format(p,spec['styles']['equation']['paragraph'])
        if inline or display:
            set_line_spacing_at_least(p, round(spec['equations']['display_min_pt' if display else 'inline_min_pt'] * 20))
    for node in doc.element.xpath('.//w:hyperlink//w:r'):
        rpr = node.get_or_add_rPr()
        for old in list(rpr.findall(qn('w:color'))):
            rpr.remove(old)
        rpr.append(element('color',val=spec['hyperlink_color']))
    apply_numbering(doc,spec)
    if spec['toc']['enabled']:
        p = doc.paragraphs[0].insert_paragraph_before() if doc.paragraphs else doc.add_paragraph()
        add_field(p, 'TOC \\o "1-' + str(spec['toc']['depth']) + '" \\h \\z \\u', '请在 Word 中更新目录域')
        settings = doc.settings.element
        for old in list(settings.findall(qn('w:updateFields'))):
            settings.remove(old)
        settings.append(element('updateFields',val='true'))


def audit_package(path, *, clean_template=False):
    with zipfile.ZipFile(path) as package:
        if len(package.infolist()) > 2000 or sum(i.file_size for i in package.infolist()) > 100_000_000:
            raise ValueError('文档包超过大小限制。')
        for name in package.namelist():
            lowered = name.lower()
            if 'vbaproject' in lowered or (clean_template and ('embeddings/' in lowered or 'activex/' in lowered)):
                raise ValueError('模板包含不允许的活动内容。')
            if name.endswith('.rels'):
                root = etree.fromstring(package.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
                for relation in root:
                    if relation.get('TargetMode') == 'External' and not relation.get('Type','').endswith('/hyperlink'):
                        raise ValueError('模板包含外部资源关系。')


def build_project(data, output: Path, report=None, cancelled=lambda: False):
    """Publish an entire new project directory only after generation succeeds."""
    spec = resolve_spec(data)
    output = output.expanduser().resolve()
    if output.exists():
        raise ValueError('模板目录已存在，请选择新的目录。')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.m2w-template-',dir=output.parent) as tmp:
        stage = Path(tmp) / 'project'
        stage.mkdir()
        dump_json(stage/'spec.json',data)
        dump_json(stage/'resolved-spec.json',spec)
        doc = Document()
        configure_document(doc,spec)
        doc.save(stage/'reference.docx')
        sample = Document()
        configure_document(sample,spec)
        sample.add_heading(spec['name'],0)
        for level,label in enumerate(('一级标题样张','二级标题样张','三级标题样张'),1):
            sample.add_heading(label,level)
        sample.add_paragraph('正文样张：中文与 English 123，用于检查字体、缩进和行距。')
        sample.add_paragraph('print("Hello Word")',style='Source Code')
        sample.add_paragraph('表 1 参数示例',style='Table Caption')
        table = sample.add_table(rows=2,cols=2)
        for cell,text in zip([c for r in table.rows for c in r.cells],['参数','数值','示例','12']):
            cell.text=text
        apply_content(sample,spec)
        sample.save(stage/'preview.docx')
        for name in ('reference.docx','preview.docx'):
            audit_package(stage/name,clean_template=True)
        dump_json(stage/'report.json',{'draft':report,'warnings':['目录页码需在 Word 中更新。'] if spec['toc']['enabled'] else [],'defaults':'缺省字段来自 m2w-default@1；完整值见 resolved-spec.json。'})
        dump_json(stage/'manifest.json',{'schema_version':'1.0','generator_version':'0.5.0','sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in stage.iterdir()}})
        if cancelled():
            raise ValueError('模板任务已取消，未发布文件。')
        os.rename(stage,output)
    return output
